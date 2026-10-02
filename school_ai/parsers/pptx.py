"""PPTX: каждый слайд — раздел «Слайд N», номер слайда сохраняется как страница."""

from pptx import Presentation

from .base import Block, ParsedDocument


def parse(path) -> ParsedDocument:
    prs = Presentation(str(path))
    blocks: list[Block] = []

    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        blocks.append(
            Block(f"Слайд {number}" + (f". {title}" if title else ""), kind="heading", heading_level=1, page=number)
        )

        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        blocks.append(Block(text, page=number))
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        blocks.append(Block(" | ".join(cells), kind="table", page=number))

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(Block(f"Заметки докладчика: {notes}", page=number))

    return ParsedDocument(blocks=blocks, pages=len(prs.slides), parser="pptx")
