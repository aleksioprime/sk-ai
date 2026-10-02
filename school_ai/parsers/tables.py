"""
XLSX и CSV.

Каждый лист — раздел. Каждая строка таблицы превращается в текст вида
«Заголовок1: значение; Заголовок2: значение», чтобы поиск понимал,
к какой колонке относится значение.
"""

import csv
from pathlib import Path

from .base import Block, ParsedDocument
from .text import _read


def _rows_to_blocks(rows: list[list], page: int | None = None) -> list[Block]:
    rows = [[("" if v is None else str(v).strip()) for v in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []

    header = rows[0]
    has_header = len(rows) > 1 and all(header) and len(set(header)) == len(header)
    blocks = []
    for row in rows[1:] if has_header else rows:
        if has_header:
            parts = [f"{h}: {v}" for h, v in zip(header, row) if v]
        else:
            parts = [v for v in row if v]
        blocks.append(Block("; ".join(parts), kind="table", page=page))
    return blocks


def parse_xlsx(path) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    for number, sheet in enumerate(wb.worksheets, start=1):
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        sheet_blocks = _rows_to_blocks(rows, page=number)
        if sheet_blocks:
            blocks.append(Block(f"Лист «{sheet.title}»", kind="heading", heading_level=1, page=number))
            blocks.extend(sheet_blocks)
    return ParsedDocument(blocks=blocks, pages=len(wb.worksheets), parser="xlsx")


def parse_csv(path) -> ParsedDocument:
    text = _read(Path(path))
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(text.splitlines(), dialect))
    return ParsedDocument(blocks=_rows_to_blocks(rows), parser="csv")
